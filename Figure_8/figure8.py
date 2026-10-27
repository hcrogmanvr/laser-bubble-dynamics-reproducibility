#!/usr/bin/env python3
"""Figure 8: authentic experimental diameter-squared growth analysis.
Usage: python figure8.py --input 'The four Graphs Fig-2 Data (2).xlsx' --output .
Requires numpy scipy matplotlib openpyxl. No synthetic model data or bar charts.
"""
import argparse, csv
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from scipy import stats
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

COLORS={'Pink':'#bd438b','Blue':'#2376b7','Green':'#2c9569','Red':'#c94a46'}
SHEETS={'Pink':'Pink G','Blue':'Blue G','Green':'Green G','Red':'Red Graph'}

def read_data(path):
    wb=load_workbook(path,data_only=True)
    out={}
    for name,sname in SHEETS.items():
        ws=wb[sname]
        rows=[]
        for row in ws.iter_rows(min_row=4,values_only=True):
            if isinstance(row[0],(int,float)) and isinstance(row[1],(int,float)):
                rows.append((float(row[0])*60,float(row[1]),row[2]))
        a=np.asarray([[r[0],r[1]] for r in rows]); a=a[np.argsort(a[:,0])]
        peak=int(np.argmax(a[:,1])); t=a[:peak+1,0]; d=a[:peak+1,1]
        fit=stats.linregress(t,d*d)
        crit=stats.t.ppf(.975,len(t)-2)
        out[name]=dict(time=a[:,0],diam=a[:,1],tfit=t,dfit=d,peak=peak,fit=fit,lower=fit.slope-crit*fit.stderr,upper=fit.slope+crit*fit.stderr)
    return out

def save(fig,out,name):
    fig.savefig(out/(name+'.png'),dpi=250,bbox_inches='tight')
    fig.savefig(out/(name+'.pdf'),bbox_inches='tight')
    plt.close(fig)

def draw(data,out):
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'figure.facecolor':'white'})
    red=data['Red']; fit=red['fit']
    fig,axs=plt.subplots(2,1,figsize=(8,9),constrained_layout=True)
    ax=axs[0]; ts=np.linspace(0,60,200)
    ax.plot(red['tfit'],red['dfit']**2,'o',c=COLORS['Red'],ms=5,label='Measured red trajectory: growth phase')
    ax.plot(ts,fit.intercept+fit.slope*ts,c='#24577b',lw=2,label=f'OLS fit: {fit.slope:.2f} µm²/s')
    # Literature range is descriptive slope-scale reference, anchored to OLS intercept; not a statistical band.
    ax.fill_between(ts,fit.intercept+30*ts,fit.intercept+80*ts,color='#adcbdc',alpha=.24,label='Published 30–80 µm²/s slope envelope (illustrative)')
    ax.set(xlim=(0,60),xlabel='Time (s)',ylabel='Squared diameter D² (µm²)'); ax.set_ylim(bottom=0); ax.legend(fontsize=8,loc='upper left'); ax.grid(alpha=.19)
    ax=axs[1]; ax.plot(red['time'],red['diam'],'o-',c=COLORS['Red'],ms=4,lw=1.2,label='Measured diameter (including contraction)')
    ax.plot(ts,np.sqrt(np.maximum(fit.intercept+fit.slope*ts,0)),c='#24577b',lw=2,label='Diameter from D² growth-phase regression')
    ax.set(xlim=(0,60),xlabel='Time (s)',ylabel='Diameter D (µm)'); ax.set_ylim(bottom=0); ax.legend(fontsize=8); ax.grid(alpha=.19)
    axs[0].text(.02,.02,'First diameter maximum at 60.54 s is included in the fit but lies outside the displayed axis.',transform=axs[0].transAxes,fontsize=8,color='#555')
    save(fig,out,'Figure_8A_Red_Growth')

    fig,axs=plt.subplots(2,1,figsize=(8.5,9),constrained_layout=True)
    for name,item in data.items():
        col=COLORS[name]; f=item['fit']; tf=item['tfit']; xx=np.linspace(max(0,tf[0]),min(60,tf[-1]),100)
        axs[0].scatter(tf, item['dfit']**2,color=col,s=26,alpha=.75)
        axs[0].plot(xx,f.intercept+f.slope*xx,c=col,lw=2,label=f'{name}: {f.slope:.2f} µm²/s')
    axs[0].set(xlim=(0,60),xlabel='Time (s)',ylabel='Squared diameter D² (µm²)');axs[0].set_ylim(bottom=0);axs[0].legend(ncol=2,fontsize=8);axs[0].grid(alpha=.19)
    names=list(data); xx=np.arange(len(names)); slopes=[data[n]['fit'].slope for n in names]
    lows=[slopes[i]-data[n]['lower'] for i,n in enumerate(names)]; ups=[data[n]['upper']-slopes[i] for i,n in enumerate(names)]
    axs[1].axhspan(30,80,color='#adcbdc',alpha=.24,label='Published descriptive range: 30–80 µm²/s')
    for i,n in enumerate(names):
        axs[1].errorbar(i,slopes[i],yerr=np.array([[lows[i]],[ups[i]]]),fmt='o',markersize=8,capsize=5,color=COLORS[n],ecolor=COLORS[n])
        axs[1].annotate(f'{slopes[i]:.2f}',(i,slopes[i]),xytext=(0,10),textcoords='offset points',ha='center',color=COLORS[n],fontsize=9)
    axs[1].set(xticks=xx,xticklabels=names,xlim=(-.5,3.5),ylim=(20,85),ylabel='Growth-phase slope d(D²)/dt (µm²/s)',xlabel='Experimental trajectory')
    axs[1].grid(axis='y',alpha=.19);axs[1].legend(fontsize=8,loc='upper left')
    save(fig,out,'Figure_8B_Four_Traces_and_Slope_CI')

    fig,ax=plt.subplots(figsize=(9,5.5),constrained_layout=True)
    for n,item in data.items(): ax.plot(item['time'],item['diam'],'o-',ms=3,lw=1.3,c=COLORS[n],label=n)
    ax.set(xlim=(0,450),xlabel='Time (s)',ylabel='Measured diameter (µm)');ax.set_ylim(bottom=0);ax.legend();ax.grid(alpha=.19)
    save(fig,out,'Figure_8C_Full_Expansion_Contraction')

def write_tables(data,out):
    heads=['Trajectory','n growth points','First growth time (s)','First maximum time (s)','First maximum diameter (µm)','Slope d(D²)/dt (µm²/s)','95% CI lower','95% CI upper','R²','Intercept (µm²)','OLS slope SE']
    records=[]
    for n,d in data.items():
        f=d['fit'];records.append([n,len(d['tfit']),d['tfit'][0],d['tfit'][-1],d['dfit'][-1],f.slope,d['lower'],d['upper'],f.rvalue**2,f.intercept,f.stderr])
    with (out/'Figure_8_Results_Table.csv').open('w',newline='') as f:
        w=csv.writer(f);w.writerow(heads);w.writerows(records)
    wb=Workbook();ws=wb.active;ws.title='Regression Summary';ws.append(['Figure 8 – Experimental growth-phase regression (measured data; not model validation)']);ws.append(heads)
    for row in records: ws.append(row)
    ws.freeze_panes='B3';ws.auto_filter.ref=f'A2:K{ws.max_row}'
    for c in ws[2]: c.fill=PatternFill('solid',fgColor='173A58');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[2].height=36
    for col in 'ABCDEFGHIJK':ws.column_dimensions[col].width=19
    ws.column_dimensions['A'].width=23;ws.column_dimensions['F'].width=29
    for row in ws.iter_rows(min_row=3):
        for c in row[1:]:c.number_format='0.0000'
    notes=wb.create_sheet('Read Me')
    for line in ['Measured workbook source: The four Graphs Fig-2 Data (2).xlsx',
        'Growth phase = first valid measurement through the first recorded maximum diameter (inclusive).',
        'Fit ordinary least squares D² = intercept + slope × time (time converted minutes to seconds).',
        '95% slope CI = slope ± t(0.975,n−2) × standard error; conditional on OLS assumptions.',
        'The workbook calls diameter units mm, but source paper Figure 2 uses µm. Calculations adopt µm provisionally.',
        'The workbook Y-error=2 is not drawn: its meaning (SD/SEM/tolerance/etc.) is unknown.',
        'Published 30–80 µm²/s is an approximate descriptive comparison, not a confidence or prediction band.',
        'Figure 8 is measured experimental calibration, NOT independent validation of a physical simulation.',
        'Source paper DOI: https://doi.org/10.1007/s00249-024-01701-3']:
        notes.append([line]);notes.column_dimensions['A'].width=120
    raw=wb.create_sheet('Source Values')
    raw.append(['Trace','Time (min)','Time (s)','Diameter (assumed µm)','Diameter² (µm²)','Growth phase?'])
    for n,d in data.items():
        for i,(t,diam) in enumerate(zip(d['time'],d['diam'])):
            j=raw.max_row+1;raw.append([n,t/60,t,diam,f'=D{j}^2',i<=d['peak']])
    for c in raw[1]: c.fill=PatternFill('solid',fgColor='173A58');c.font=Font(color='FFFFFF',bold=True)
    for col in 'ABCDEF':raw.column_dimensions[col].width=24
    wb.save(out/'Figure_8_Results_Table.xlsx')
    return records

def main():
    p=argparse.ArgumentParser();p.add_argument('--input',required=True);p.add_argument('--output',default='.');a=p.parse_args()
    out=Path(a.output);out.mkdir(parents=True,exist_ok=True)
    data=read_data(a.input);draw(data,out);rows=write_tables(data,out)
    for r in rows: print(r[0],r[1],round(r[5],4),round(r[6],4),round(r[7],4),round(r[8],5))
if __name__=='__main__': main()
